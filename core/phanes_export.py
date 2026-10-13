"""Opt-in import of user-obtained Phanes .xlsx leaderboard exports.

Phanes does NOT offer a public API and forbids automated bot queries.
This module never queries Phanes or Telegram. It treats uploaded workbooks as
untrusted third-party data, not authenticated call histories or trade fills.
"""
from __future__ import annotations
import hashlib, io, math, re, zipfile
from datetime import date, datetime
from openpyxl import load_workbook

MAX_XLSX_BYTES = 1_500_000
MAX_ZIP_UNCOMPRESSED = 12_000_000
MAX_CALLS = 1500
MAX_COLUMNS = 40
MAX_CELL_LENGTH = 220


def clean_cell(v) -> str:
    if v is None:
        return ''
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if isinstance(v, float) and not math.isfinite(v):
            return ''
        return str(v)
    text = re.sub(r'[\x00-\x08\x0b-\x1f]', '', str(v))
    return text[:MAX_CELL_LENGTH]


def workbook_to_archive(raw: bytes, filename: str) -> dict:
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('Only .xlsx Phanes leaderboard exports are supported')
    if not raw or len(raw)>MAX_XLSX_BYTES:
        raise ValueError('Workbook empty or exceeds 1.5MB limit')
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise ValueError('Invalid XLSX/ZIP workbook')
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        members=z.infolist()
        if len(members)>200 or sum(f.file_size for f in members)>MAX_ZIP_UNCOMPRESSED:
            raise ValueError('Workbook exceeds safe decompression limits')
        if any(f.flag_bits & 1 for f in members):
            raise ValueError('Encrypted workbooks not supported')
    try:
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        try:
            names=wb.sheetnames
            calls_name=next((n for n in names if n.strip().lower()=='calls'),None)
            if not calls_name:
                raise ValueError('Phanes export must contain a Calls worksheet')
            ws=wb[calls_name]
            if ws.max_column and ws.max_column>MAX_COLUMNS:
                raise ValueError('Calls worksheet has too many columns')
            iterator=ws.iter_rows(values_only=True)
            header=None
            for ix,row in enumerate(iterator):
                if ix>=12: break
                values=[clean_cell(v).strip() for v in row[:MAX_COLUMNS]]
                if sum(bool(v) for v in values)>=2 and any(
                    k in ' '.join(v.lower() for v in values) for k in ('caller','token','entry','ath','contract','market cap','ca','user','time','date')
                ):
                    header=values
                    break
            if not header:
                raise ValueError('Could not recognize a Phanes Calls worksheet header (send a redacted sample for mapping)')
            columns=[]
            for i,name in enumerate(header):
                base=name or f'column_{i+1}'
                field=base
                count=1
                while field in columns:
                    count+=1
                    field=f'{base}_{count}'
                columns.append(field[:80])
            rows=[]
            for row in iterator:
                if len(rows)>=MAX_CALLS:
                    raise ValueError('Over 1,500 Calls rows; export a shorter period')
                d={columns[i]:clean_cell(v) for i,v in enumerate(row[:len(columns)]) if i<len(columns) and v is not None}
                if not any(d.values()):continue
                rows.append(d)
            if not rows:raise ValueError('Calls worksheet contains no call rows')
            digest=hashlib.sha256(raw).hexdigest()
            calls=[{'source':'phanes_xlsx_user_upload','file_sha256':digest,'row_number':i+1,'fields':r}
                   for i,r in enumerate(rows)]
            extras={n:wb[n].max_row for n in names if n.strip().lower() in ('summary','top callers')}
            return {'schema':'talos.external.phanes_excel.v1','status':'PARSED_UNVERIFIED',
                    'file_sha256':digest,'filename':filename[:160],'columns':columns,'calls':calls,
                    'call_count':len(calls),'sheet_names':names[:30],'other_sheet_rows_estimated':extras,
                    'provider_authenticated':False,'publisher_authenticated':False,
                    'call_time_verified':False,'price_execution_verified':False,
                    'live_orders_enabled':False,'notifications_sent':False,
                    'warning':'User-supplied Phanes workbook. ATH/current multiples are not realized, executable trade returns. Verify original posting timestamps and liquidity separately.'}
        finally:wb.close()
    except ValueError:raise
    except Exception as e:
        raise ValueError('Cannot parse this workbook as a supported XLSX file') from e
